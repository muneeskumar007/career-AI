"""
Raw text extraction from resume files. Kept deliberately separate from
structured parsing (skills/education/experience — that's Module 3's
resume_parser.py) so this module has exactly one job and is trivial to unit
test: bytes in, plain text out.
"""

import io
import logging

import docx
import fitz  # PyMuPDF

from app.utils.exceptions import AppException

logger = logging.getLogger(__name__)


class TextExtractionError(AppException):
    status_code = 422
    message = "Couldn't extract text from this file"


def extract_text_from_pdf(content: bytes) -> str:
    try:
        with fitz.open(stream=content, filetype="pdf") as pdf:
            pages = [page.get_text() for page in pdf]
        text = "\n".join(pages).strip()
    except Exception as exc:  # PyMuPDF raises its own exception types per failure mode
        logger.warning("PDF text extraction failed: %s", exc)
        raise TextExtractionError() from exc

    if not text:
        raise TextExtractionError(
            "No extractable text found — this may be a scanned/image-only PDF, which isn't supported yet"
        )
    return text


def extract_text_from_docx(content: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        # Resumes often use tables for skills/contact layout — paragraphs alone miss that text.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        text = "\n".join(parts).strip()
    except Exception as exc:
        logger.warning("DOCX text extraction failed: %s", exc)
        raise TextExtractionError() from exc

    if not text:
        raise TextExtractionError("No extractable text found in this document")
    return text


def extract_text(content: bytes, extension: str) -> str:
    if extension == "pdf":
        return extract_text_from_pdf(content)
    if extension == "docx":
        return extract_text_from_docx(content)
    raise TextExtractionError(f"Unsupported extension for extraction: {extension}")
